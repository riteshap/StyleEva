
import numpy as np
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from matplotlib.ticker import MultipleLocator

def fun_4():
    task1 = {"embedding":{"same":[0.43, 0.43, 0.41, 0.39, 0.35], "higher":[0.66,0.71,0.71,0.71,0.72], "lower":[0.69,0.74,0.74,0.74,0.74]},
             "classification":{"same":[0.72,0.63,0.62,0.59, None], "higher":[0.79,0.83,0.86,0.86, None], "lower":[0.79,0.84,0.87,0.86, None]},
             "pairwise":{"same":[0.73,0.70,0.66,0.66, 0.59], "higher":[0.82,0.88,0.89,0.88,0.87], "lower":[0.82,0.88,0.88,0.88,0.87]}}
    task8 = {"embedding":{"same":[0.44,0.46,0.44,0.43,0.41], "higher":[0.67,0.72,0.74,0.75,0.75], "lower":[0.70,0.76,0.77,0.78,0.78]},
             "classification":{"same":[0.63,0.56,0.50,0.45, None], "higher":[0.68,0.75,0.76,0.77, None], "lower":[0.68,0.75,0.77,0.78, None]},
             "pairwise":{"same":[0.66,0.60,0.55,0.52,0.44], "higher":[0.72,0.80,0.81,0.80,0.82], "lower":[0.72,0.80,0.82,0.79,0.82]}}


    x_labels = ["Q3", "Q5", "Q7", "Q9", "Dense"]

    model_titles = {
        "embedding": "Generic Embedding",
        "classification": "Multiclass Classification",
        "pairwise": "Pairwise Relational Learning"
    }

    relation_styles = {
        "same":   {"marker": "o", "linestyle": "-"},
        "higher": {"marker": "s", "linestyle": "--"},
        "lower":  {"marker": "^", "linestyle": "-."}
    }


    def plot_relation_f1(task_data, dataset_name, save_path=None):
        fig, axes = plt.subplots(
            1, 3,
            figsize=(13.5, 4),
            sharey=True
        )

        x = np.arange(len(x_labels))

        for ax, model in zip(
            axes,
            ["embedding", "classification", "pairwise"]
        ):
            for relation in ["same", "higher", "lower"]:
                values = [
                    np.nan if v is None else v
                    for v in task_data[model][relation]
                ]

                ax.plot(
                    x,
                    values,
                    label=relation.capitalize(),
                    linewidth=2,
                    markersize=6,
                    **relation_styles[relation]
                )

            ax.set_title(model_titles[model], fontsize=11)
            ax.set_xticks(x)
            ax.set_xticklabels(x_labels)
            ax.set_xlabel("Granularity")
            ax.set_ylim(0.2, 1.0)


            ax.set_yticks(np.arange(0.2, 1.01, 0.1))


            ax.grid(
                axis="y",
                linestyle="--",
                linewidth=0.8,
                alpha=0.4
            )


            ax.spines["top"].set_visible(False)
            ax.spines["right"].set_visible(False)

        axes[0].set_ylabel("F1 Score")


        handles, labels = axes[0].get_legend_handles_labels()
        fig.legend(
            handles,
            labels,
            loc="upper center",
            ncol=3,
            frameon=False,
            bbox_to_anchor=(0.5, 1.03)
        )

        fig.suptitle(dataset_name, fontsize=13, y=1.10)
        fig.tight_layout()

        if save_path:
            plt.savefig(
                save_path,
                dpi=300,
                bbox_inches="tight"
            )

        plt.show()


    plot_relation_f1(
        task1,
        "Yelp Sentiment",
        "yelp_relation_f1.png"
    )

    plot_relation_f1(
        task8,
        "Shakespeare",
        "shakespeare_relation_f1.png"
    )


def plot_threshold_results(
    yelp_result,
    shakespeare_result,
    model_name,
    save_path=None
):
    datasets = [
        ("Yelp Sentiment", yelp_result),
        ("Shakespeare", shakespeare_result)
    ]


    granularities = ["Dense", "Q9", "Q7", "Q5", "Q3"]


    line_styles = ["-", "--", "-.", ":", "-"]
    markers = ["o", "s", "^", "D", "v"]

    fig, axes = plt.subplots(
        1, 2,
        figsize=(11, 4),
        sharey=True
    )

    for ax, (dataset_name, result) in zip(axes, datasets):


        data = np.array(result, dtype=float)


        thresholds = data[:, 0]


        for i, granularity in enumerate(granularities):
            ax.plot(
                thresholds,
                data[:, i + 1],
                label=granularity,
                linewidth=2,
                linestyle=line_styles[i],
                marker=markers[i],
                markersize=4,
                markevery=max(1, len(thresholds) // 10)
            )


        ax.set_title(dataset_name, fontsize=12)


        ax.set_xlabel("Threshold")
        ax.xaxis.set_major_locator(MultipleLocator(0.1))


        ax.set_ylim(0.15, 0.75)
        ax.yaxis.set_major_locator(MultipleLocator(0.1))


        ax.grid(
            axis="y",
            linestyle="--",
            linewidth=0.8,
            alpha=0.4
        )


        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)

        ax.spines["left"].set_linewidth(0.8)
        ax.spines["bottom"].set_linewidth(0.8)


    axes[0].set_ylabel("Macro-F1")


    handles, labels = axes[0].get_legend_handles_labels()

    fig.legend(
        handles,
        labels,
        loc="upper center",
        ncol=5,
        frameon=False,
        bbox_to_anchor=(0.5, 1.03)
    )


    fig.suptitle(
        model_name,
        fontsize=13,
        y=1.10
    )

    fig.tight_layout()

    if save_path is not None:
        plt.savefig(
            save_path,
            dpi=300,
            bbox_inches="tight"
        )

    plt.show()


def fun_6():

    wb = load_workbook("4.6result.xlsx", data_only=True)

    print(wb.sheetnames)

    ws = wb["task8-s"]

    # print(ws["A2"].value)
    # print(ws["B2"].value)


    ws = wb["task1-s"]
    yelp_result = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cur_list = []
        for cell in row:
            cur_list.append(cell)
        yelp_result.append(cur_list)

    ws = wb["task8-s"]
    shakespeare_result = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cur_list = []
        for cell in row:
            cur_list.append(cell)
        shakespeare_result.append(cur_list)
    pass

    plot_threshold_results(
        yelp_result,
        shakespeare_result,
        model_name="Generic Embedding",
        save_path="embedding_threshold_analysis.png"
    )

if __name__ == "__main__":
    fun_6()